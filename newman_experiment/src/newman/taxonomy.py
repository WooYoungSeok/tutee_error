"""Newman taxonomy: 16 adopted source error types with their Newman stage, name and definition; 10 excluded types.

* Labels in the source data are resolved per dataset through the listed source aliases only (plan 3.3-1/2): EIC
  `unit_conversion_error` (adopted) and Stepwise `Unit conversion error` (excluded) never merge.
* The stage is a deterministic function of the type (plan 3.3-6): N = mapping(E). `condition()` is the only way a
  (N, E) pair is formed, so a stage can never be attached to a type it does not belong to.
* `verify_workbook` re-reads the user's workbook (sheet "Newman 재분류") and compares every row with this file: the
  name (column A, exact text), the definition (column B, exact text; empty = no definition) and the decision
  (column D "영석 분류": a stage or 제외). Column C (혁규 분류) is never applied.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

from .common import hash_obj, rel, sha256_file

STAGE_IDS = ("reading", "comprehension", "transformation", "process_skills")
EXPECTED_ADOPTED = 16
EXPECTED_EXCLUDED = 10
EXPECTED_PER_STAGE = {"reading": 2, "comprehension": 2, "transformation": 6, "process_skills": 6}
UNIT_RELATED_TYPES = frozenset({"mathedu.measurement_error", "eic.unit_conversion_error"})  # plan 5.5 (user confirmed)
EXCLUDED = "excluded"


class TaxonomyError(RuntimeError):
    pass


@dataclass(frozen=True)
class Stage:
    id: str
    name: str
    definition: str


@dataclass(frozen=True)
class ErrorType:
    id: str
    dataset: str
    name: str                    # workbook column A: shown to the Student and the verifiers
    source_aliases: tuple[str, ...]
    newman_stage: str
    definition: str | None       # workbook column B; None = the workbook gives none (the line is left out)
    unit_related: bool


@dataclass(frozen=True)
class ExcludedType:
    id: str
    dataset: str
    name: str
    source_aliases: tuple[str, ...]


def flat(text: Any) -> str:
    """Whitespace-collapsed text; case is kept (EIC 'Unit Conversion Error' != Stepwise 'Unit conversion error')."""
    return re.sub(r"\s+", " ", str(text)).strip() if text is not None else ""


def find_file(path: Path) -> Path | None:
    """The path, or a file in its directory whose name is equal after Unicode NFC (uploads from macOS are NFD)."""
    if path.exists():
        return path
    if not path.parent.is_dir():
        return None
    want = unicodedata.normalize("NFC", path.name)
    hits = [p for p in path.parent.iterdir() if unicodedata.normalize("NFC", p.name) == want]
    return hits[0] if len(hits) == 1 else None


class Taxonomy:
    def __init__(self, raw: Mapping[str, Any], path: Path | None = None):
        self.raw = raw
        self.path = path
        self.sha256 = sha256_file(path) if path else hash_obj(raw)
        self.version = str(raw.get("version", ""))
        self.mapping_source = dict(raw.get("mapping_source") or {})
        self.framework = flat(raw.get("framework"))  # Newman's Error Analysis overview shown before the stage
        self.stages = {sid: Stage(sid, s["name"], flat(s["definition"])) for sid, s in (raw.get("stages") or {}).items()}
        self.datasets = list(raw.get("datasets") or [])
        self.types: dict[str, ErrorType] = {}
        for t in raw.get("error_types") or []:
            et = ErrorType(id=t["id"], dataset=t["dataset"], name=t["name"], source_aliases=tuple(t["source_aliases"]),
                           newman_stage=t["newman_stage"], definition=flat(t["definition"]) or None,
                           unit_related=bool(t.get("unit_related", False)))
            if et.id in self.types:
                raise TaxonomyError(f"duplicate error type id {et.id}")
            self.types[et.id] = et
        self.excluded: dict[str, ExcludedType] = {}
        for t in raw.get("excluded_types") or []:
            ex = ExcludedType(id=t["id"], dataset=t["dataset"], name=t["name"], source_aliases=tuple(t["source_aliases"]))
            if ex.id in self.excluded or ex.id in self.types:
                raise TaxonomyError(f"duplicate type id {ex.id}")
            self.excluded[ex.id] = ex
        self._alias: dict[tuple[str, str], str] = {}
        for item in self.all_items():
            for alias in item.source_aliases:
                key = (item.dataset, alias)
                if key in self._alias:
                    raise TaxonomyError(f"alias {alias!r} of {item.dataset} used by {self._alias[key]} and {item.id}")
                self._alias[key] = item.id
        self.validate()

    @classmethod
    def load(cls, path: str | Path) -> "Taxonomy":
        import yaml

        path = Path(path)
        return cls(yaml.safe_load(path.read_text(encoding="utf-8")), path)

    def all_items(self) -> list[ErrorType | ExcludedType]:
        return [*self.types.values(), *self.excluded.values()]

    # --- checks ---------------------------------------------------------------

    def validate(self) -> None:
        problems = []
        if not self.framework:
            problems.append("framework (the Newman's Error Analysis overview) is missing")
        if tuple(self.stages) != STAGE_IDS:
            problems.append(f"stages must be exactly {STAGE_IDS} (no Encoding type is adopted), got {tuple(self.stages)}")
        for s in self.stages.values():
            if not s.name or not s.definition:
                problems.append(f"stage {s.id} needs a name and a definition")
        if len(self.types) != EXPECTED_ADOPTED:
            problems.append(f"{len(self.types)} adopted types, expected {EXPECTED_ADOPTED}")
        if len(self.excluded) != EXPECTED_EXCLUDED:
            problems.append(f"{len(self.excluded)} excluded types, expected {EXPECTED_EXCLUDED}")
        per_stage = {sid: sum(1 for t in self.types.values() if t.newman_stage == sid) for sid in STAGE_IDS}
        if per_stage != EXPECTED_PER_STAGE:
            problems.append(f"adopted types per stage {per_stage}, expected {EXPECTED_PER_STAGE}")
        names: dict[str, str] = {}
        for item in self.all_items():
            if item.dataset not in self.datasets:
                problems.append(f"{item.id}: unknown dataset {item.dataset}")
            if not item.id.startswith(item.dataset + "."):
                problems.append(f"{item.id}: id must start with its dataset")
            if not item.source_aliases or not flat(item.name):
                problems.append(f"{item.id}: needs a name and a source alias")
            if flat(item.name) in names:  # workbook rows are matched by exact name
                problems.append(f"name {item.name!r} used by {names[flat(item.name)]} and {item.id}")
            names[flat(item.name)] = item.id
        for t in self.types.values():
            if t.newman_stage not in STAGE_IDS:
                problems.append(f"{t.id}: stage {t.newman_stage!r}")
        unit = {t.id for t in self.types.values() if t.unit_related}
        if unit != UNIT_RELATED_TYPES:
            problems.append(f"unit_related types {sorted(unit)}, the plan restricts exactly {sorted(UNIT_RELATED_TYPES)}")
        if problems:
            raise TaxonomyError("taxonomy is inconsistent:\n  - " + "\n  - ".join(problems))

    # --- lookups ----------------------------------------------------------------

    def resolve(self, dataset: str, raw_label: str) -> tuple[str, str | None]:
        """('adopted', type id) | ('excluded', type id) | ('unknown', None). Exact alias match inside the dataset."""
        type_id = self._alias.get((dataset, raw_label))
        if type_id is None:
            return "unknown", None
        return ("adopted" if type_id in self.types else EXCLUDED), type_id

    def stage_of(self, type_id: str) -> str:
        return self.types[type_id].newman_stage

    def types_for_dataset(self, dataset: str) -> list[ErrorType]:
        return sorted((t for t in self.types.values() if t.dataset == dataset), key=lambda t: t.id)

    def adopted_ids(self) -> list[str]:
        return sorted(self.types)

    def condition(self, type_id: str) -> dict[str, str | None]:
        """The (N, E) condition shown to the Student and the verifiers; N always follows E."""
        t = self.types[type_id]
        s = self.stages[t.newman_stage]
        return {"stage_id": s.id, "stage_name": s.name, "stage_definition": s.definition,
                "error_id": t.id, "error_name": t.name, "error_definition": t.definition}

    def prompt_values(self, type_id: str) -> dict[str, str | None]:
        """Template values; source_error_definition None = the prompt line holding it is left out."""
        c = self.condition(type_id)
        return {"newman_framework": self.framework,
                "newman_stage_name": c["stage_name"], "newman_stage_definition": c["stage_definition"],
                "source_error_name": c["error_name"], "source_error_definition": c["error_definition"]}

    def summary(self) -> dict[str, Any]:
        return {"version": self.version, "sha256": self.sha256, "path": rel(self.path) if self.path else None,
                "adopted": {t.id: t.newman_stage for t in sorted(self.types.values(), key=lambda t: t.id)},
                "excluded": sorted(self.excluded), "unit_related": sorted(UNIT_RELATED_TYPES),
                "without_definition": sorted(t.id for t in self.types.values() if t.definition is None)}


# --- mapping workbook -------------------------------------------------------------------


def parse_decision(cell: Any, taxonomy: Taxonomy) -> tuple[str, str]:
    """A column-D cell -> (stage id or 'excluded', trailing note). 'Transformation(9)' -> ('transformation', '(9)')."""
    text = flat(cell)
    prefix = str(taxonomy.mapping_source.get("excluded_prefix", "제외"))
    if text.startswith(prefix):
        return EXCLUDED, text[len(prefix):].strip()
    low = re.sub(r"[\s_-]+", " ", text).casefold()
    for sid, stage in taxonomy.stages.items():
        name = re.sub(r"[\s_-]+", " ", stage.name).casefold()
        if low.startswith(name) and (len(low) == len(name) or not low[len(name)].isalpha()):
            return sid, low[len(name):].strip()
    raise TaxonomyError(f"column D value {text!r} is neither a stage name nor '{prefix}...'")


def read_workbook(path: Path, taxonomy: Taxonomy) -> dict[str, Any]:
    """Rows below the header row (the row whose decision column holds decision_header) with a label."""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    src = taxonomy.mapping_source
    wb = openpyxl.load_workbook(path, data_only=True)
    if src["sheet"] not in wb.sheetnames:
        raise TaxonomyError(f"sheet {src['sheet']!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[src["sheet"]]
    col = {k: column_index_from_string(src[k]) for k in ("label_column", "definition_column", "decision_column", "other_annotator_column")}
    header_row = next((r for r in range(1, min(ws.max_row, 30) + 1)
                       if flat(ws.cell(r, col["decision_column"]).value) == src["decision_header"]), None)
    if header_row is None:
        preview = [[ws.cell(r, c).value for c in range(1, 7)] for r in range(1, min(ws.max_row, 4) + 1)]
        raise TaxonomyError(f"no header {src['decision_header']!r} in column {src['decision_column']} of {src['sheet']!r} "
                            f"(first rows: {preview}); set mapping_source columns in configs/taxonomy.yaml")
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(r, col["label_column"]).value
        if label is None or not flat(label):
            continue
        rows.append({"row": r, "label": flat(label), "definition": flat(ws.cell(r, col["definition_column"]).value) or None,
                     "decision_cell": flat(ws.cell(r, col["decision_column"]).value),
                     "other_annotator_cell": ws.cell(r, col["other_annotator_column"]).value})
    return {"header_row": header_row, "headers": {k: ws.cell(header_row, c).value for k, c in col.items()}, "rows": rows}


def verify_workbook(path: Path, taxonomy: Taxonomy) -> dict[str, Any]:
    """Compare every workbook row with the taxonomy. `problems` empty = the workbook confirms this file."""
    src = taxonomy.mapping_source
    report: dict[str, Any] = {"verified": False, "workbook": path.name, "sha256": sha256_file(path),
                              "expected_sha256": src.get("sha256"), "sheet": src.get("sheet"),
                              "columns": {k: src.get(k) for k in ("label_column", "definition_column", "decision_column",
                                                                   "other_annotator_column")},
                              "types": {}, "unmatched_rows": [], "problems": []}
    problems = report["problems"]
    if report["sha256"] != src.get("sha256"):
        problems.append(f"workbook sha256 {report['sha256']} != expected {src.get('sha256')}")
    book = read_workbook(path, taxonomy)
    report.update(header_row=book["header_row"], headers=book["headers"])
    if "혁규" not in str(book["headers"].get("other_annotator_column") or ""):
        problems.append(f"column {src['other_annotator_column']} header is {book['headers'].get('other_annotator_column')!r}, "
                        "expected the 혁규 분류 column")
    by_name = {flat(item.name): item for item in taxonomy.all_items()}
    for row in book["rows"]:
        item = by_name.get(row["label"])
        if item is None:
            report["unmatched_rows"].append(row)
            continue
        try:
            decision, note = parse_decision(row["decision_cell"], taxonomy)
        except TaxonomyError as exc:
            problems.append(f"row {row['row']}: {exc}")
            continue
        if item.id in report["types"]:
            problems.append(f"{item.id} appears in rows {report['types'][item.id]['row']} and {row['row']}")
        expected = EXCLUDED if item.id in taxonomy.excluded else taxonomy.types[item.id].newman_stage
        if decision != expected:
            problems.append(f"row {row['row']} {item.id}: workbook says {decision!r}, taxonomy says {expected!r}")
        if item.id in taxonomy.types and row["definition"] != taxonomy.types[item.id].definition:
            problems.append(f"row {row['row']} {item.id}: definition differs from workbook column {src['definition_column']}")
        report["types"][item.id] = {"row": row["row"], "label_cell": row["label"], "definition_cell": row["definition"],
                                    "mapping_cell_raw": row["decision_cell"], "decision": decision, "note": note}
    missing = sorted({i.id for i in taxonomy.all_items()} - set(report["types"]))
    if missing:
        problems.append(f"taxonomy types not found in the workbook: {missing}")
    if report["unmatched_rows"]:
        r0 = report["unmatched_rows"][0]
        problems.append(f"{len(report['unmatched_rows'])} workbook rows match no taxonomy name (e.g. row {r0['row']}: {r0['label']!r})")
    report["verified"] = not problems
    return report


def unverified_mapping_report(taxonomy: Taxonomy, reason: str) -> dict[str, Any]:
    """Stand-in for smoke/audit runs without the workbook; real runs refuse data made with it."""
    return {"verified": False, "reason": reason, "workbook": None, "sha256": None,
            "expected_sha256": taxonomy.mapping_source.get("sha256"),
            "types": {i.id: {"row": None, "mapping_cell_raw": None,
                             "decision": EXCLUDED if i.id in taxonomy.excluded else taxonomy.types[i.id].newman_stage}
                      for i in sorted(taxonomy.all_items(), key=lambda i: i.id)},
            "unmatched_rows": [], "problems": [reason]}
