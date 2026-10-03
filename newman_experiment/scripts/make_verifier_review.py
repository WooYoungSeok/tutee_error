#!/usr/bin/env python3
"""Blind human-review sheet for the verifier-A vs gpt-5.6-sol disagreement on the API baselines' RL-test outputs
(user request 2026-10-02, after rl/EXPERIMENTS.md N12-4).

Rows (incorrect solutions only, seed 42, split evenly over the two generator models and spread over stages):
  50 A rejected / gpt-5.6-sol accepted (the disagreement), 10 both accepted, 10 both rejected (controls).
The rows are shuffled and the review sheet hides both verdicts and the category; the `key` sheet holds them.

Usage (from newman_experiment/, `source env.sh`):  python scripts/make_verifier_review.py
Output: reports/review_verifier_disagreement_20261002.xlsx (sheets review, key, guide) + .csv (review sheet only)
"""

from __future__ import annotations

import csv
import json
import random
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from newman.common import read_jsonl, resolve  # noqa: E402
from newman.taxonomy import Taxonomy  # noqa: E402

GENERATORS = ("gpt-5.6-sol", "gpt-5.1")
QUOTA = {"A0_sol1": 25, "A1_sol1": 5, "A0_sol0": 5}  # per generator
OUT = "reports/review_verifier_disagreement_20261002"


def passes(labels) -> bool:
    return bool(labels) and all(x == "aligned" for x in labels)


def spread_pick(items, n, rng):
    """n items spread over stages (round robin over shuffled per-stage lists)."""
    by = defaultdict(list)
    for it in items:
        by[it["newman_stage"]].append(it)
    for v in by.values():
        rng.shuffle(v)
    out, stages = [], sorted(by)
    while len(out) < n and any(by.values()):
        for s in stages:
            if by[s] and len(out) < n:
                out.append(by[s].pop())
    return out


def main() -> int:
    rng = random.Random(42)
    tax = Taxonomy.load(resolve("configs/taxonomy.yaml"))
    test = {r["condition_id"]: r for r in read_jsonl(resolve("data/prepared/rl/test.jsonl"))}
    ref = {r["condition_id"]: r["reference_answer"] for r in read_jsonl(resolve("data/prepared/rl/privileged.jsonl"))}
    rows = []
    for g in GENERATORS:
        scored = read_jsonl(resolve(f"outputs/api_baselines_prelim_verifierA/test_eval/api_{g}/rollouts/step_000000.jsonl"))
        sol = defaultdict(list)
        for v in read_jsonl(resolve(f"outputs/api_baselines_verifier_gpt-5.6-sol/api_{g}/verdicts_raw.jsonl")):
            sol[(v["condition_id"], v["k"])].append(v["prediction"])
        cats = defaultdict(list)
        for r in scored:
            if r["answer_check"]["verdict"] != "incorrect":
                continue
            a_labels, s_labels = r["verifier_b"]["labels"], sol[(r["condition_id"], r["k"])]  # verifier_b = A server (N12-3)
            cats[f"A{int(passes(a_labels))}_sol{int(passes(s_labels))}"].append({**r, "generator": g, "a_labels": a_labels, "sol_labels": s_labels})
        for cat, n in QUOTA.items():
            for it in spread_pick(cats[cat], n, rng):
                rows.append({**it, "category": cat})
    rng.shuffle(rows)
    review, key = [], []
    for i, r in enumerate(rows, 1):
        pv = tax.prompt_values(r["source_error_id"])
        rid = f"R{i:02d}"
        review.append({"review_id": rid, "stage": pv["newman_stage_name"], "stage_definition": pv["newman_stage_definition"],
                       "error_type": pv["source_error_name"], "error_type_definition": pv.get("source_error_definition") or "(no definition)",
                       "question": test[r["condition_id"]]["question"], "reference_answer": ref[r["condition_id"]],
                       "extracted_answer": r["answer_check"]["extracted_answer"], "solution": r["solution"],
                       "human_judgment (aligned / not_aligned / unsure)": "", "first_wrong_step (copy or describe)": "", "note": ""})
        key.append({"review_id": rid, "category": r["category"], "generator": r["generator"], "condition_id": r["condition_id"],
                    "verifier_A_labels": " ".join(r["a_labels"]), "gpt-5.6-sol_labels": " ".join(r["sol_labels"])})

    out = resolve(OUT)
    with open(f"{out}.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(review[0]))
        w.writeheader()
        w.writerows(review)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(list(review[0]))
    for r in review:
        ws.append(list(r.values()))
    widths = {"A": 9, "B": 14, "C": 45, "D": 22, "E": 40, "F": 50, "G": 10, "H": 10, "I": 70, "J": 18, "K": 30, "L": 25}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for col in ("J", "K", "L"):
        for c in ws[col][1:]:
            c.fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"aligned,not_aligned,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(review) + 1}")
    ws.freeze_panes = "B2"

    ks = wb.create_sheet("key")
    ks.append(list(key[0]))
    for r in key:
        ks.append(list(r.values()))
    ks.sheet_state = "hidden"

    gs = wb.create_sheet("guide")
    for line in (
        "검수 방법",
        "1. review 시트에서 문제, 지정 단계·유형(정의 포함), 풀이를 읽고, 풀이에서 '처음 틀린 단계'의 오류가 지정 단계와 유형에 맞는지 판정한다.",
        "2. 판정은 verifier 지시문과 같은 기준: 단계와 유형이 모두 맞아야 aligned. 최종 답이 틀렸다는 것만으로는 aligned가 아니다.",
        "3. human_judgment 칸에 aligned / not_aligned / unsure 중 하나, first_wrong_step 칸에 처음 틀린 단계를 적는다.",
        "4. 모든 행은 gpt-5-nano가 오답으로 판정한 풀이다. verifier 판정과 표본 종류는 숨겨져 있다(key 시트, 숨김). 검수를 마친 뒤에 연다.",
        f"표본: {len(review)}행 = A 거부·gpt-5.6-sol 통과 50, 둘 다 통과 10, 둘 다 거부 10 (생성 모델 gpt-5.6-sol / gpt-5.1 반반, seed 42, 섞음)",
    ):
        gs.append([line])
    gs.column_dimensions["A"].width = 140
    wb.move_sheet("guide", offset=-2)
    wb.active = 1
    wb.save(f"{out}.xlsx")
    counts = defaultdict(int)
    for k in key:
        counts[(k["generator"], k["category"])] += 1
    print(f"{len(review)} rows -> {OUT}.xlsx / .csv", dict(counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
